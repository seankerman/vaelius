"""Faithful original documents and canonical native passages.

Objects preserve original bytes. PostgreSQL stores upload states and source
references; parsed passages reuse the existing knowledge pipeline without models.
"""
from __future__ import annotations

from dataclasses import asdict, dataclass
from html.parser import HTMLParser
import hashlib
import io
import json
from pathlib import Path
import re
import tempfile
import time
from urllib.parse import quote
import zipfile

from agentclient.enterprise_capture import normalize_capture, redact
from agenthub.source_objects import FILE_LIMIT, READ_SIZE, ObjectConflict, ObjectCorrupt, ObjectMissing, bounded_spool, object_key, verify


@dataclass(frozen=True)
class Location:
    start: int
    end: int
    section: str = ''
    page: int | None = None
    kind: str = 'paragraph'


@dataclass(frozen=True)
class PreparedDocument:
    status: str
    text: str
    media_type: str
    classification: str
    locations: tuple[Location, ...] = ()


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []; self.skip = 0; self.heading = 0

    def handle_starttag(self, tag, attrs):
        if tag in {'script', 'style', 'template', 'noscript'}:
            self.skip += 1
        if self.skip:
            return
        if re.fullmatch(r'h[1-6]', tag):
            self.parts.append('\n' + '#' * int(tag[1]) + ' '); self.heading += 1
        elif tag in {'p', 'div', 'section', 'article', 'tr', 'li', 'br'}:
            self.parts.append('\n')
        elif tag in {'td', 'th'}:
            self.parts.append('\t')

    def handle_endtag(self, tag):
        if tag in {'script', 'style', 'template', 'noscript'}:
            self.skip = max(0, self.skip - 1)
        elif not self.skip and re.fullmatch(r'h[1-6]', tag):
            self.parts.append('\n'); self.heading = max(0, self.heading - 1)
        elif not self.skip and tag in {'p', 'div', 'section', 'article', 'tr', 'li'}:
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def _text_locations(text):
    from agenthub.local_connectors import native_sections
    return tuple(Location(start, end, heading) for start, end, heading in
        native_sections(text, markdown=True))


def prepare_document(raw, filename, *, media_type=None, expanded_limit=100 * 1024 * 1024):
    """Parser failures never destroy or gate authorized original retrieval."""
    suffix = Path(filename).suffix.lower()
    media = media_type or {'.txt':'text/plain', '.md':'text/markdown',
        '.html':'text/html', '.htm':'text/html', '.pdf':'application/pdf',
        '.docx':'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        }.get(suffix, 'application/octet-stream')
    try:
        if suffix in {'.txt', '.md'}:
            text = raw.decode('utf-8-sig')
            return PreparedDocument('indexed', text, media, 'document', _text_locations(text))
        if suffix in {'.html', '.htm'}:
            parser = _HTMLText(); parser.feed(raw.decode('utf-8-sig')); parser.close()
            text = re.sub(r'\n{3,}', '\n\n', ''.join(parser.parts)).strip()
            return PreparedDocument('indexed', text, media, 'document', _text_locations(text))
        if suffix == '.pdf':
            from pypdf import PdfReader
            pdf = PdfReader(io.BytesIO(raw), strict=True)
            if pdf.is_encrypted:
                return PreparedDocument('encrypted', '', media, 'document')
            text = ''; locations = []
            for number, page in enumerate(pdf.pages, 1):
                start = len(text); text += (page.extract_text() or '') + '\n'
                locations.append(Location(start, len(text), f'Page {number}', number, 'page'))
            if not text.strip():
                return PreparedDocument('ocr_required', '', media, 'document')
            return PreparedDocument('indexed', text, media, 'document', tuple(locations))
        if suffix == '.docx':
            # Bound decompression before letting the maintained parser open it.
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(item.file_size for item in archive.infolist()) > expanded_limit:
                    return PreparedDocument('expanded_bound', '', media, 'document')
            from docx import Document
            from docx.table import Table
            from docx.text.paragraph import Paragraph
            doc = Document(io.BytesIO(raw)); parts = []
            for block in doc.iter_inner_content():
                if isinstance(block, Paragraph):
                    style = block.style.name if block.style else ''
                    match = re.fullmatch(r'Heading ([1-6])', style)
                    parts.append(('#' * int(match[1]) + ' ' if match else '') + block.text + '\n')
                elif isinstance(block, Table):
                    parts.extend('\t'.join(cell.text for cell in row.cells) + '\n' for row in block.rows)
            text = ''.join(parts)
            return PreparedDocument('indexed', text, media, 'document', _text_locations(text))
    except (ImportError, ModuleNotFoundError):
        return PreparedDocument('parser_unavailable', '', media, 'document')
    except Exception:
        # A parser traceback may contain source material, so retain only a
        # content-free disposition. Operators can retry a new parser version.
        return PreparedDocument('parser_failed', '', media, 'document')
    return PreparedDocument('unsupported', '', media, 'document')


def download_headers(filename, length, sha256):
    safe = Path(filename.replace('\\', '/')).name
    safe = re.sub(r'[\x00-\x1f\x7f]', '_', safe)[:180] or 'document'
    # Force attachment/octet-stream even for untrusted HTML and SVG originals.
    ascii_name = re.sub(r'[^A-Za-z0-9._ -]', '_', safe)
    return {'Content-Type':'application/octet-stream', 'Content-Length':str(length),
        'Content-Disposition':f'attachment; filename="{ascii_name}"; filename*=UTF-8\'\'{quote(safe)}',
        'X-Content-Type-Options':'nosniff', 'Cache-Control':'no-store',
        'X-Source-SHA256':sha256}


class DocumentStore:
    """Original-document service over the single enterprise authority.

    Static migration 003_objects.sql is applied by the PostgreSQL profile. Store
    transactions use its existing psycopg-backed canonical State connection.
    """
    def __init__(self, store, objects, *, file_limit=FILE_LIMIT):
        self.store = store; self.objects = objects; self.file_limit = file_limit

    def _connection(self, ctx, ident, *, ingest=True):
        with self.store.open() as state:
            return dict(self.store._connection(state.db, ctx, ident, ingest=ingest))

    def _metrics(self,identity,source_id):
        # Metrics describe one accepted source revision. A retry repairs a
        # partial metrics write using the same stable IDs, never charges again.
        with self.store.open() as state:
            if getattr(state.db,'dialect',None)!='postgres' or not state.db.execute(
                    "SELECT to_regclass('public.cloud_metrics')").fetchone()[0]:
                return
            row=state.db.execute('SELECT byte_length,parser_status FROM backend_document_versions WHERE source_id=%s',
                (source_id,)).fetchone()
            if not row:return
            passages=state.db.execute('SELECT count(*) FROM backend_native_artifacts WHERE source_id=%s',
                (source_id,)).fetchone()[0]
        from agenthub.cloud_ops import Meter
        meter=Meter(self.store);details={'source_revision':source_id,'status':row['parser_status']}
        for kind,amount in {'original_bytes':row['byte_length'],'accepted_source_revisions':1,
                'parsing_attempts':1,'native_passages':passages}.items():
            meter.metric('document:'+identity+':'+kind,kind,amount,details=details)

    def ingest(self, ctx, connection, external_id, version, filename, stream, *,
               title=None, media_type=None, source_url='', occurred_at='unknown'):
        self.store._need(ctx, 'ingest')
        enrolled = self._connection(ctx, connection)
        if 'document' not in json.loads(enrolled['source_types']):
            raise ValueError('document_connection_required')
        if not re.fullmatch(r'[1-9][0-9]{0,14}', str(version)):
            raise ValueError('invalid_document_version')
        if not external_id or len(external_id) > 256 or not filename or len(filename) > 256:
            raise ValueError('invalid_document_identity')
        metadata = {'external_id': external_id, 'filename': filename, 'title': title or filename,
                    'source_url': source_url, 'media_type': media_type or ''}
        if redact(metadata)[0] != metadata:
            raise ValueError('secret_document_metadata_rejected')
        with bounded_spool(stream, limit=self.file_limit) as (spool, length, sha):
            # Reject recognizable credentials before storing original bytes. Do
            # not silently alter originals: their hashes and locations must agree.
            raw = spool.read(self.file_limit + 1); spool.seek(0)
            prepared = prepare_document(raw, filename, media_type=media_type)
            try:
                original_text = raw.decode('utf-8-sig')
            except UnicodeDecodeError:
                original_text = ''
            if any(redact(text)[0] != text for text in (original_text, prepared.text)):
                raise ValueError('secret_document_rejected')
            key = object_key(ctx['tenant'], connection + ':' + external_id, str(version), sha)
            identity = hashlib.sha256(json.dumps([ctx['tenant'], connection, external_id,
                str(version)], separators=(',', ':')).encode()).hexdigest()
            now = time.time()
            # Identity conflicts are checked before object writes. Upload state
            # commits before crossing the database/object-store boundary.
            with self.store.open() as state, state.db:
                prior = state.db.execute('SELECT * FROM backend_object_uploads WHERE id=%s FOR UPDATE', (identity,)).fetchone()
                if prior and (prior['sha256'] != sha or prior['byte_length'] != length):
                    raise ObjectConflict('document_version_conflict')
                if not prior:
                    state.db.execute('''INSERT INTO backend_object_uploads
                        (id,tenant,connection,external_id,version,object_key,sha256,byte_length,status,updated)
                        VALUES(%s,%s,%s,%s,%s,%s,%s,%s,'uploading',%s)
                        ON CONFLICT(id) DO NOTHING''',
                        (identity,ctx['tenant'],connection,external_id,str(version),key,sha,length,now))
                prior = state.db.execute('SELECT * FROM backend_object_uploads WHERE id=%s FOR UPDATE', (identity,)).fetchone()
                if prior['sha256'] != sha or prior['byte_length'] != length:
                    raise ObjectConflict('document_version_conflict')
                if prior['status'] == 'accepted':
                    retained = dict(prior)
                else:
                    retained = None
                    state.db.execute("UPDATE backend_object_uploads SET status='uploading',updated=%s WHERE id=%s", (now,identity))
            if retained:
                verify(self.objects.head(key), length, sha)
                self._metrics(identity,retained['source_id'])
                self.project_temporal(ctx,retained['source_id'])
                return self.describe(ctx, retained['source_id']) | {'disposition':'duplicate'}
            info = self.objects.put(key, spool, expected_sha256=sha)
            verify(info, length, sha)
            verify(self.objects.head(key), length, sha)
            with self.store.open() as state, state.db:
                state.db.execute("UPDATE backend_object_uploads SET status='uploaded',updated=%s WHERE id=%s", (time.time(),identity))
        # Source intake is a bounded attachment reference. Large source files do
        # not weaken or overflow the agent-event contract.
        event = normalize_capture({'hook_event_name':'NativeSource','event_id':external_id,
            'revision':str(version),'timestamp':occurred_at,'source_url':source_url},
            enrolled['project'],connection,source_type='document',origin='original_document')
        event['event']['role'] = 'source'; event['disposition'] = 'accepted'
        event['blocks'] = [{'type':'attachment','value':{'reference':'source-object:'+identity,
            'media_type':prepared.media_type,'name':filename}}]
        event, _ = redact(event)
        result = self.store.ingest_general(ctx, event); source_id = result['source_id']
        indexed = 0
        if prepared.status == 'indexed' and prepared.text.strip():
            from agenthub.local_connectors import Connector
            safe, redactions = redact(prepared.text)
            indexed = Connector(self.store)._native(source_id, {'source_type':'document',
                'blocks':[{'type':'text','value':safe}], 'title':title or filename,
                'id':external_id,'source_url':source_url})
        with self.store.open() as state, state.db:
            state.db.execute('''INSERT INTO backend_document_versions
                (source_id,tenant,connection,external_id,version,object_key,sha256,byte_length,
                 filename,title,media_type,parser_status,locations,original_disposition,created)
                VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'original',%s)
                ON CONFLICT(source_id) DO NOTHING''',
                (source_id,ctx['tenant'],connection,external_id,str(version),key,sha,length,
                 filename,title or filename,prepared.media_type,prepared.status,
                 json.dumps([asdict(x) for x in prepared.locations]),time.time()))
            state.db.execute("UPDATE backend_object_uploads SET status='accepted',source_id=%s,updated=%s WHERE id=%s", (source_id,time.time(),identity))
            self._temporal_passages(state.db,source_id,title or filename)
        self._metrics(identity,source_id)
        return self.describe(ctx,source_id) | {'disposition':result['disposition'],'passages':indexed,'model_calls':0}

    @staticmethod
    def _temporal_passages(db,source_id,title,retained_source_authorizer=None):
        from agenthub.processing.temporal import admit_native_assertion
        rows=db.execute('''SELECT d.active_revision_id FROM backend_native_artifacts n
            JOIN knowledge_documents d ON d.document_id=n.document_id
            WHERE n.source_id=%s AND n.artifact_kind='native_document' AND d.lifecycle='active' ''',(source_id,)).fetchall()
        for row in rows:admit_native_assertion(db,row['active_revision_id'],source_id,subject=title,
            retained_source_authorizer=retained_source_authorizer)
        return len(rows)

    def project_temporal(self,ctx,source_id):
        """Authorized idempotent derivative repair, preserving originals/heads."""
        with self.store.delivery_lock(),self.store.open() as state,state.db:
            item=self._authorize(state.db,ctx,source_id)
            def authorized(ident):
                self._authorize(state.db,ctx,ident)
                return True
            count=self._temporal_passages(state.db,item['source_id'],item['title'],authorized)
            self._authorize(state.db,ctx,item['source_id'])
        return {'source_id':item['source_id'],'version':item['version'],'passages_checked':count,'model_calls':0}

    def _authorize(self, db, ctx, source_id, version=None):
        from agenthub.enterprise import Denied
        self.store._need(ctx, 'read')
        selected = db.execute('SELECT * FROM backend_document_versions WHERE source_id=%s AND tenant=%s', (source_id,ctx['tenant'])).fetchone()
        if not selected:
            raise Denied()
        if version is not None:
            selected = db.execute('''SELECT * FROM backend_document_versions WHERE tenant=%s
                AND connection=%s AND external_id=%s AND version=%s''',
                (ctx['tenant'],selected['connection'],selected['external_id'],str(version))).fetchone()
            if not selected:
                raise Denied()
        # Historical versions use the current source head and policy, not their
        # original access snapshot or now-retired source.active flag.
        current = db.execute('''SELECT s.* FROM backend_source_heads h JOIN enterprise_sources s
            ON s.id=h.source_id WHERE h.connection=%s AND h.external_id=%s''',
            (selected['connection'],selected['external_id'])).fetchone()
        if not self.store._visible_source(db,ctx,current):
            raise Denied()
        if db.execute('''SELECT 1 FROM enterprise_deletion_journal WHERE tenant=%s
                AND source_id=%s AND operation IN ('withdraw','delete','correct')''',
                (ctx['tenant'],selected['source_id'])).fetchone():
            raise Denied()
        historical = db.execute('SELECT * FROM enterprise_sources WHERE id=%s', (selected['source_id'],)).fetchone()
        if not historical or not self.store._visible_source(db,ctx,dict(historical)|{'active':1}):
            # Revision retirement alone is permitted; explicit historical policy
            # narrowing and lifecycle denials remain effective independently.
            raise Denied()
        upload = db.execute('SELECT status FROM backend_object_uploads WHERE object_key=%s', (selected['object_key'],)).fetchone()
        if not upload or upload['status'] != 'accepted':
            raise ObjectMissing('document_not_fully_accepted')
        return dict(selected)

    def describe(self, ctx, source_id, *, version=None):
        with self.store.open() as state:
            item = self._authorize(state.db,ctx,source_id,version)
        return {k:item[k] for k in ('source_id','version','filename','title','media_type',
            'parser_status','byte_length','sha256','original_disposition')} | {
            'download_path':'/enterprise/v3/source-documents/download',
            'download_method':'POST',
            'download_request':{'source_id':item['source_id'],'version':item['version']},
            'requires_authenticated_download':True}

    def list(self, ctx, *, title='', limit=20):
        if type(limit) is not int or not 1 <= limit <= 100 or len(title) > 256:
            raise ValueError('invalid_document_listing')
        found = []
        with self.store.open() as state:
            self.store._need(ctx,'read')
            if hasattr(self.store,'current_identity'):self.store.current_identity(state.db,ctx)
            rows = state.db.execute('''SELECT v.source_id FROM backend_document_versions v
                JOIN enterprise_sources s ON s.id=v.source_id
                WHERE v.tenant=%s AND (%s='' OR position(lower(%s) in lower(v.title))>0
                    OR position(lower(%s) in lower(v.filename))>0)
                AND s.tenant=%s AND (s.visibility!='private' OR s.owner=%s)
                ORDER BY v.title,v.source_id''', (ctx['tenant'],title,title,title,ctx['tenant'],ctx['actor']))
            from agenthub.enterprise import Denied
            for row in rows:
                try:
                    item = self._authorize(state.db,ctx,row['source_id'])
                except Denied:
                    continue
                found.append({k:item[k] for k in ('source_id','version','filename','title','parser_status')})
                if len(found) == limit:
                    break
        return found

    def fetch(self, ctx, source_id, *, version=None):
        """Return a verified private temporary binary stream and safe headers.

        API must perform its current-authorization delivery check again immediately
        before headers/bytes are sent. No bearer URL or server path is returned.
        """
        with self.store.open() as state:
            item = self._authorize(state.db,ctx,source_id,version)
        spool = tempfile.TemporaryFile()
        try:
            with self.objects.open(item['object_key']) as source, bounded_spool(source,
                    limit=self.file_limit,expected_sha256=item['sha256']) as (verified,length,sha):
                if length != item['byte_length']:
                    raise ObjectCorrupt('retained_source_length_mismatch')
                while block := verified.read(READ_SIZE):
                    spool.write(block)
            spool.seek(0)
            with self.store.open() as state:
                self._authorize(state.db,ctx,item['source_id'],item['version'])
            return spool,download_headers(item['filename'],length,sha)
        except Exception:
            spool.close(); raise

    def cleanup(self, key, *, older_than):
        """Remove only unreferenced, inactive orphan reservations after a grace."""
        with self.store.open() as state, state.db:
            rows = state.db.execute('SELECT * FROM backend_object_uploads WHERE object_key=%s FOR UPDATE', (key,)).fetchall()
            if any(row['status'] in {'accepted','uploading'} or row['updated'] > older_than for row in rows):
                raise ValueError('object_referenced_or_inflight')
            if state.db.execute('SELECT 1 FROM backend_document_versions WHERE object_key=%s', (key,)).fetchone():
                raise ValueError('object_referenced_or_inflight')
            # The row lock serializes recovery claims and cleanup for this key.
            self.objects.delete(key)
            state.db.execute("UPDATE backend_object_uploads SET status='removed',updated=%s WHERE object_key=%s", (time.time(),key))
