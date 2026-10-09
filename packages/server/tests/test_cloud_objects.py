import hashlib
import io
import json
import os
from pathlib import Path
import tempfile
import unittest
import uuid

from agenthub.source_objects import (FILE_LIMIT, FileSourceObjects, S3SourceObjects,
    ObjectCorrupt, ObjectMissing, bounded_spool, object_key)
from agenthub.document_ingest import DocumentStore, prepare_document, download_headers

FIXTURE_PATH = Path(__file__).resolve().parents[1]/'tests/fixtures/service/cloud_objects_v1.json'
FIXTURE = json.loads(FIXTURE_PATH.read_text())


def pdf_bytes(text):
    # A deterministic synthetic text PDF; no copyrighted/private input.
    escaped = text.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
    stream = f'BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET'.encode()
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
        b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
    body = b'%PDF-1.4\n'; offsets = [0]
    for number, value in enumerate(objects,1):
        offsets.append(len(body));body += str(number).encode()+b' 0 obj\n'+value+b'\nendobj\n'
    xref = len(body)
    body += b'xref\n0 6\n0000000000 65535 f \n'+b''.join(f'{offset:010} 00000 n \n'.encode() for offset in offsets[1:])
    return body+b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n'+str(xref).encode()+b'\n%%EOF\n'


def docx_bytes():
    from docx import Document
    fixture = next(x for x in FIXTURE['documents'] if x['format']=='docx')
    document = Document();document.add_heading(fixture['heading'],level=1)
    document.add_paragraph(fixture['text']);table = document.add_table(rows=2,cols=2)
    for i,row in enumerate(fixture['table']):
        for j,value in enumerate(row):
            table.cell(i,j).text=value
    out=io.BytesIO();document.save(out);return out.getvalue()


class ObjectContractTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.objects=FileSourceObjects(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def exercise(self, adapter):
        raw=b'Original synthetic bytes\x00\xff';sha=hashlib.sha256(raw).hexdigest()
        key=object_key('a','guide','1',sha)
        receipt=adapter.put(key,io.BytesIO(raw),expected_sha256=sha)
        self.assertEqual((receipt.length,receipt.sha256),(len(raw),sha))
        self.assertEqual(adapter.put(key,io.BytesIO(raw)),receipt)
        with adapter.open(key) as result:
            self.assertEqual(result.read(),raw)
        self.assertEqual(adapter.head(key),receipt)
        self.assertNotEqual(key,object_key('b','guide','1',sha))
        with self.assertRaises(ObjectCorrupt):
            adapter.put(key,io.BytesIO(b'changed'))
        adapter.delete(key)
        with self.assertRaises(ObjectMissing):
            adapter.open(key)

    def test_file_contract_and_restart(self):
        self.exercise(self.objects)
        raw=b'durable';sha=hashlib.sha256(raw).hexdigest();key=object_key('a','s','1',sha)
        self.objects.put(key,io.BytesIO(raw))
        self.assertEqual(FileSourceObjects(self.temp.name).head(key).sha256,sha)

    def test_bounded_reads_and_exact_file_ceiling(self):
        class FiniteStream:
            def __init__(self,size):self.left=size;self.calls=[]
            def read(self,n):
                if n<0:raise AssertionError('unbounded read')
                self.calls.append(n);block=b'x'*min(self.left,n);self.left-=len(block);return block
        stream=FiniteStream(FILE_LIMIT)
        with bounded_spool(stream) as (_,length,_):self.assertEqual(length,FILE_LIMIT)
        self.assertLessEqual(max(stream.calls),65536)
        with self.assertRaisesRegex(ValueError,'source_file_too_large'):
            with bounded_spool(FiniteStream(FILE_LIMIT+1)):pass

    def test_paths_and_cloud_endpoint_fail_closed(self):
        for bad in ('../../escape','https://example.com/file','sources/a','/etc/passwd'):
            with self.assertRaises(ValueError):self.objects.open(bad)
        for endpoint in ('https://s3.amazonaws.com','https://not-local.example','http://user:pass@localhost:1'):
            with self.assertRaises(ValueError):S3SourceObjects(endpoint,'bucket','fixture','fixture')

    def test_corrupt_file_verified(self):
        raw=b'source';sha=hashlib.sha256(raw).hexdigest();key=object_key('a','s','1',sha)
        self.objects.put(key,io.BytesIO(raw))
        self.objects._path(key).write_bytes(b'corrupted')
        self.assertNotEqual(self.objects.head(key).sha256,sha)


class MotoObjectContractTests(unittest.TestCase):
    def test_real_sdk_against_local_moto_protocol(self):
        from moto_test_fixture import sdk_objects
        with sdk_objects('synthetic-originals') as adapter:
            ObjectContractTests.exercise(self,adapter)


class DocumentParserTests(unittest.TestCase):
    def test_frozen_manifest_and_faithful_markdown(self):
        manifest=json.loads(FIXTURE_PATH.with_name('cloud_objects_v1_manifest.json').read_text())
        self.assertEqual(hashlib.sha256(FIXTURE_PATH.read_bytes()).hexdigest(),manifest['sha256'])
        guide=FIXTURE['documents'][0];result=prepare_document(guide['v1'].encode(),'guide.md')
        self.assertEqual(result.text,guide['v1']);self.assertEqual(result.status,'indexed')
        self.assertEqual(len([x for x in result.locations if x.section.startswith('Section')]),20)
        for location in result.locations:
            self.assertTrue(result.text[location.start:location.end].strip())

    def test_html_preserves_table_and_removes_scripts_from_text_only(self):
        fixture=FIXTURE['documents'][1];result=prepare_document(fixture['body'].encode(),'table.html')
        for expected in fixture['expected']:self.assertIn(expected,result.text)
        for excluded in fixture['excluded']:self.assertNotIn(excluded,result.text)
        self.assertIn('Owner\tFormat',result.text)

    def test_text_pdf_and_docx_keep_locations(self):
        text=next(x['text'] for x in FIXTURE['documents'] if x['format']=='pdf' and 'text' in x)
        pdf=prepare_document(pdf_bytes(text),'guide.pdf')
        self.assertEqual(pdf.status,'indexed');self.assertIn(text,pdf.text);self.assertEqual(pdf.locations[0].page,1)
        word=prepare_document(docx_bytes(),'guide.docx')
        self.assertEqual(word.status,'indexed');self.assertIn('Owner\tDataset',word.text)
        self.assertIn('Alice\tmaple.csv',word.text);self.assertEqual(word.locations[0].section,'Maple Archive')

    def test_unsupported_broken_encrypted_and_ocr_are_explicit(self):
        self.assertEqual(prepare_document(b'\x00\xff','data.bin').status,'unsupported')
        self.assertEqual(prepare_document(b'not a PDF','broken.pdf').status,'parser_failed')
        from pypdf import PdfReader,PdfWriter
        writer=PdfWriter();writer.append_pages_from_reader(PdfReader(io.BytesIO(pdf_bytes('secret'))));writer.encrypt('fixture-only')
        out=io.BytesIO();writer.write(out)
        self.assertEqual(prepare_document(out.getvalue(),'encrypted.pdf').status,'encrypted')
        writer=PdfWriter();writer.add_blank_page(width=200,height=200);out=io.BytesIO();writer.write(out)
        self.assertEqual(prepare_document(out.getvalue(),'scan.pdf').status,'ocr_required')

    def test_safe_original_headers(self):
        headers=download_headers('../../evil\r\nSet-Cookie: bad.html',123,'a'*64)
        self.assertEqual(headers['Content-Type'],'application/octet-stream')
        self.assertNotIn('\r',headers['Content-Disposition']);self.assertNotIn('\n',headers['Content-Disposition'])
        self.assertEqual(headers['Cache-Control'],'no-store')


@unittest.skipUnless(os.environ.get('CLOUD_TEST_DSN'),'real PostgreSQL fixture DSN required')
class DocumentPostgresTests(unittest.TestCase):
    def setUp(self):
        from agenthub.postgres import PostgresEnterpriseStore
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        tenant=os.environ.get('CLOUD_TEST_TENANT','acme');suffix=uuid.uuid4().hex[:10]
        self.store=PostgresEnterpriseStore(self.root/'authority',os.environ['CLOUD_TEST_DSN'],tenant)
        self.store.create_organization(tenant);self.project='doc-'+suffix
        self.alice='alice-'+suffix;self.bob='bob-'+suffix
        self.store.create_project(tenant,self.project)
        for actor in (self.alice,self.bob):
            self.store.create_principal(tenant,actor);self.store.set_membership(tenant,self.project,actor,True)
        actions=['ingest','read','source_read','policy','withdraw','correct']
        self.ctx={actor:self.store.authenticate(self.store.enroll(tenant,actor,actor,actions)) for actor in (self.alice,self.bob)}
        self.connection='objects-'+suffix
        self.store.enroll_connection(self.ctx[self.alice],self.connection,'documents-'+suffix,self.project,['document'],visibility='team',reader_ids=[self.alice,self.bob])
        self.objects=FileSourceObjects(self.root/'originals');self.documents=DocumentStore(self.store,self.objects)

    def tearDown(self):self.temp.cleanup()

    def ingest(self,raw,version=1,filename='guide.md',external='guide'):
        return self.documents.ingest(self.ctx[self.alice],self.connection,external,str(version),filename,io.BytesIO(raw),title='Maple Guide')

    def test_twenty_passages_single_original_and_exact_versions_current_acl(self):
        fixture=FIXTURE['documents'][0];raw=fixture['v1'].encode();first=self.ingest(raw)
        self.assertGreaterEqual(first['passages'],20)
        self.assertEqual(len(list((self.root/'originals').iterdir())),1)
        stream,headers=self.documents.fetch(self.ctx[self.bob],first['source_id'])
        with stream:self.assertEqual(stream.read(),raw)
        second=self.ingest(fixture['v2'].encode(),2)
        stream,_=self.documents.fetch(self.ctx[self.bob],second['source_id'],version='1')
        with stream:self.assertEqual(stream.read(),raw)
        self.assertEqual({item['version'] for item in self.documents.list(self.ctx[self.bob],title='Maple Guide')},{'1','2'})
        from agenthub.enterprise import Denied
        wrong=self.ctx[self.bob]|{'tenant':'wrong-tenant'}
        with self.assertRaises(Denied):self.documents.describe(wrong,first['source_id'])
        self.store.connection_policy(self.ctx[self.alice],self.connection,reader_ids=[self.alice])
        for source in (first,second):
            with self.assertRaises(Denied):self.documents.fetch(self.ctx[self.bob],source['source_id'])
        self.assertFalse(self.documents.list(self.ctx[self.bob]))

    def test_parser_independent_fetch_idempotency_and_conflict(self):
        raw=b'\x00\xffunsupported source';first=self.ingest(raw,filename='original.bin')
        self.assertEqual(first['parser_status'],'unsupported');self.assertEqual(first['passages'],0)
        self.assertEqual(self.ingest(raw,filename='original.bin')['disposition'],'duplicate')
        from agenthub.source_objects import ObjectConflict
        with self.assertRaises(ObjectConflict):self.ingest(b'changed',filename='original.bin')
        stream,_=self.documents.fetch(self.ctx[self.bob],first['source_id'])
        with stream:self.assertEqual(stream.read(),raw)

    def test_upload_success_metadata_failure_recovery_and_corrupt_missing_objects(self):
        raw=b'Original retained after metadata failure.'
        original=self.store.ingest_general
        self.store.ingest_general=lambda *a,**k: (_ for _ in ()).throw(RuntimeError('synthetic database fault'))
        with self.assertRaises(RuntimeError):self.ingest(raw)
        self.store.ingest_general=original
        self.assertEqual(len(list((self.root/'originals').iterdir())),1)
        result=self.ingest(raw)
        with self.store.open() as state:
            row=state.db.execute('SELECT object_key FROM backend_document_versions WHERE source_id=%s',(result['source_id'],)).fetchone()
        with self.assertRaises(ValueError):self.documents.cleanup(row['object_key'],older_than=float('inf'))
        self.objects._path(row['object_key']).write_bytes(b'corrupt')
        with self.assertRaises(ObjectCorrupt):self.documents.fetch(self.ctx[self.bob],result['source_id'])
        self.objects.delete(row['object_key'])
        with self.assertRaises(ObjectMissing):self.documents.fetch(self.ctx[self.bob],result['source_id'])

    def test_zero_models_and_no_sqlite_authority(self):
        result=self.ingest(b'# Source\nMeaningful original fact.')
        self.assertEqual(result['model_calls'],0)
        self.assertFalse(list((self.root/'authority').rglob('*.sqlite')))
        with self.store.open() as state:
            self.assertEqual(state.db.execute('SELECT count(*) FROM backend_document_versions WHERE connection=%s',(self.connection,)).fetchone()[0],1)

    def test_source_revision_metrics_are_stable_across_accepted_retries(self):
        with self.store.open() as state:
            if not state.db.execute("SELECT to_regclass('public.cloud_metrics')").fetchone()[0]:
                self.skipTest('operations migration required for document metrics')
        raw=b'# Metric fixture\nMeaningful original fact.\n'
        first=self.ingest(raw);duplicate=self.ingest(raw)
        self.assertEqual(duplicate['disposition'],'duplicate')
        with self.store.open() as state:
            rows=state.db.execute('SELECT kind,amount,details FROM cloud_metrics WHERE details::jsonb->>\'source_revision\'=%s',(first['source_id'],)).fetchall()
        metrics={row['kind']:row['amount'] for row in rows}
        self.assertEqual(metrics,{'original_bytes':len(raw),'accepted_source_revisions':1,
            'parsing_attempts':1,'native_passages':first['passages']})
        self.assertTrue(all(json.loads(row['details'])=={'source_revision':first['source_id'],'status':'indexed'} for row in rows))

    def test_authorize_before_limit_finds_original_after_thousand_forbidden_metadata_rows(self):
        path=FIXTURE_PATH.with_name('cloud_document_listing_v1.json')
        manifest=json.loads(path.with_name('cloud_document_listing_v1_manifest.json').read_text())
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),manifest['sha256'])
        fixture=json.loads(path.read_text())
        allowed=self.documents.ingest(self.ctx[self.alice],self.connection,'listing-original','1','original.bin',
            io.BytesIO(b'Original metadata fixture bytes'),title=fixture['allowed_title'])
        with self.store.open() as state,state.db:
            source=dict(state.db.execute('SELECT * FROM enterprise_sources WHERE id=%s',(allowed['source_id'],)).fetchone())
            document=dict(state.db.execute('SELECT * FROM backend_document_versions WHERE source_id=%s',(allowed['source_id'],)).fetchone())
            # Bulk metadata fixture isolates discovery policy/limit behavior. It
            # is not ordinary ingestion/object storage throughput evidence.
            rows=[];versions=[];heads=[]
            for index in range(fixture['forbidden_original_metadata_rows']):
                ident=allowed['source_id']+':denied:'+str(index);external='listing-denied-'+str(index)
                denied=source|{'id':ident,'external_id':external,'payload_hash':hashlib.sha256(ident.encode()).hexdigest(),'visibility':'private'}
                metadata=document|{'source_id':ident,'external_id':external,'title':f'A{index:04d} Needle private original'}
                rows.append(tuple(denied.values()));versions.append(tuple(metadata.values()));heads.append((self.connection,external,'1',ident))
            columns=','.join(source);doc_columns=','.join(document)
            state.db.executemany('INSERT INTO enterprise_sources ('+columns+') VALUES('+','.join(['%s']*len(source))+')',rows)
            state.db.executemany('INSERT INTO backend_document_versions ('+doc_columns+') VALUES('+','.join(['%s']*len(document))+')',versions)
            state.db.executemany('INSERT INTO backend_source_heads(connection,external_id,revision,source_id) VALUES(%s,%s,%s,%s)',heads)
        result=self.documents.list(self.ctx[self.bob],title=fixture['query'],limit=fixture['limit'])
        self.assertEqual([row['source_id'] for row in result],[allowed['source_id']])
        self.assertEqual(result[0]['title'],fixture['allowed_title'])

    def test_historical_version_withdrawal_and_private_policy_are_not_bypassed(self):
        first=self.ingest(b'Original v1 bytes.')
        second=self.ingest(b'Current v2 bytes.',2)
        self.store.lifecycle(self.ctx[self.alice],{'version':'enterprise-local-1','operation':'policy',
            'target_id':first['source_id'],'expected_revision':'1','idempotency_key':'private-'+first['source_id'],
            'reason':'narrow historical original','visibility':'private'})
        from agenthub.enterprise import Denied
        with self.assertRaises(Denied):self.documents.fetch(self.ctx[self.bob],second['source_id'],version='1')
        stream,_=self.documents.fetch(self.ctx[self.bob],second['source_id'])
        with stream:self.assertEqual(stream.read(),b'Current v2 bytes.')
        self.store.lifecycle(self.ctx[self.alice],{'version':'enterprise-local-1','operation':'withdraw',
            'target_id':first['source_id'],'expected_revision':'2','idempotency_key':'withdraw-'+first['source_id'],
            'reason':'withdraw historical original'})
        with self.assertRaises(Denied):self.documents.fetch(self.ctx[self.alice],second['source_id'],version='1')

    def test_historical_originals_are_discoverable_under_current_and_historical_policy(self):
        first=self.ingest(b'Original narrow bracket.',version=1)
        second=self.ingest(b'Revised wide bracket.',version=2)
        listed=self.documents.list(self.ctx[self.bob],title='Maple Guide')
        self.assertEqual({row['version'] for row in listed},{'1','2'})
        self.assertEqual({row['source_id'] for row in listed},{first['source_id'],second['source_id']})
        stream,_=self.documents.fetch(self.ctx[self.bob],first['source_id'],version='1')
        with stream:self.assertEqual(stream.read(),b'Original narrow bracket.')
        self.store.connection_policy(self.ctx[self.alice],self.connection,reader_ids=[self.alice])
        self.assertEqual(self.documents.list(self.ctx[self.bob],title='Maple Guide'),[])


if __name__=='__main__':unittest.main()
